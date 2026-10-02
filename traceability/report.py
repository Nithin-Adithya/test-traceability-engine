"""Excel export of the traceability matrix and gap report."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from traceability.engine import TraceResult

_HEADER_FILL = PatternFill("solid", fgColor="1F3864")
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_GAP_FILL = PatternFill("solid", fgColor="FCE4D6")
_MAX_WIDTH = 80


def _write_sheet(sheet: Worksheet, headers: list[str], rows: list[list], flag_column: int | None = None) -> None:
    sheet.append(headers)
    for cell in sheet[1]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
    for row in rows:
        sheet.append(row)
        if flag_column is not None and row[flag_column] == "Uncovered":
            for cell in sheet[sheet.max_row]:
                cell.fill = _GAP_FILL
    for i, header in enumerate(headers):
        longest = max([len(header)] + [len(str(row[i])) for row in rows if row[i] is not None])
        column = sheet.cell(row=1, column=i + 1).column_letter
        sheet.column_dimensions[column].width = min(longest + 2, _MAX_WIDTH)
        if longest + 2 > _MAX_WIDTH:
            for cell in sheet[column][1:]:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
    sheet.freeze_panes = "A2"


def build_workbook(result: TraceResult, evaluation: dict | None = None) -> Workbook:
    workbook = Workbook()
    summary = result.summary()

    sheet = workbook.active
    sheet.title = "Summary"
    rows = [
        ["Requirements", summary["requirements"]],
        ["Test cases", summary["tests"]],
        ["Links", summary["links"]],
        ["Covered requirements", summary["covered_requirements"]],
        ["Uncovered requirements", summary["uncovered_requirements"]],
        ["Orphan test cases", summary["orphan_tests"]],
        ["Requirement coverage %", summary["coverage_pct"]],
        ["Similarity threshold", summary["config"]["threshold"]],
        ["Runner-up margin", summary["config"]["margin"]],
        ["Processing time (s)", summary["timings_s"].get("total")],
    ]
    if evaluation:
        rows += [
            ["Match accuracy vs ground truth", evaluation["match_accuracy"]],
            ["Link precision", evaluation["links"]["precision"]],
            ["Link recall", evaluation["links"]["recall"]],
        ]
    _write_sheet(sheet, ["Metric", "Value"], rows)

    _write_sheet(
        workbook.create_sheet("Traceability Matrix"),
        ["Requirement ID", "Requirement", "Status", "Test Count", "Covering Tests", "Best Score"],
        [[r["requirement_id"], r["requirement"], r["status"], r["test_count"], r["tests"], r["best_score"]]
         for r in result.matrix()],
        flag_column=2,
    )

    test_text = {t.id: t.embedding_text for t in result.tests}
    req_text = {r.id: r.embedding_text for r in result.requirements}
    _write_sheet(
        workbook.create_sheet("Links"),
        ["Requirement ID", "Requirement", "Test ID", "Test Case", "Score"],
        [[l.requirement_id, req_text[l.requirement_id], l.test_id, test_text[l.test_id], round(l.score, 3)]
         for l in sorted(result.links, key=lambda l: (l.requirement_id, -l.score))],
    )

    _write_sheet(
        workbook.create_sheet("Uncovered Requirements"),
        ["Requirement ID", "Requirement", "Closest Test ID", "Closest Test", "Closest Score"],
        [[r["requirement_id"], r["requirement"], r["closest_test_id"], r["closest_test"], r["closest_score"]]
         for r in result.uncovered()],
    )

    _write_sheet(
        workbook.create_sheet("Orphan Tests"),
        ["Test ID", "Test Case", "Closest Requirement ID", "Closest Requirement", "Closest Score"],
        [[r["test_id"], r["test"], r["closest_requirement_id"], r["closest_requirement"], r["closest_score"]]
         for r in result.orphans()],
    )
    return workbook


def workbook_bytes(result: TraceResult, evaluation: dict | None = None) -> bytes:
    buffer = BytesIO()
    build_workbook(result, evaluation).save(buffer)
    return buffer.getvalue()


def write_report(result: TraceResult, path: str | Path, evaluation: dict | None = None) -> Path:
    path = Path(path)
    path.write_bytes(workbook_bytes(result, evaluation))
    return path
