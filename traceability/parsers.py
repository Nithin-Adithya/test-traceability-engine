"""Readers for requirement documents and test-case repositories.

Requirements: .docx (tables, or "REQ-001: text" paragraphs), .xlsx, .csv
Tests:        .xlsx, .csv, .docx, .json, pytest .py files, or a directory of any of these
"""

from __future__ import annotations

import ast
import csv
import json
import re
from pathlib import Path
from typing import Iterable

from traceability.models import Requirement, TestCase


class ParseError(ValueError):
    """Raised when a file cannot be interpreted as requirements or test cases."""


# Column headers are matched after lower-casing and stripping non-alphanumerics.
REQUIREMENT_COLUMNS = {
    "id": {"id", "reqid", "requirementid", "reqno", "requirementno", "req", "identifier", "key", "ref"},
    "title": {"title", "name", "summary", "requirementtitle", "requirementname"},
    "text": {"description", "requirement", "requirementdescription", "requirementtext", "text",
             "statement", "details", "requirementstatement"},
}
TEST_COLUMNS = {
    "id": {"id", "testid", "testcaseid", "tcid", "caseid", "testno", "testcaseno", "identifier", "key"},
    "title": {"title", "name", "summary", "testname", "testcase", "testcasename", "testtitle",
              "testcasetitle", "scenario"},
    "text": {"description", "objective", "preconditions", "precondition", "steps", "teststeps",
             "procedure", "expected", "expectedresult", "expectedresults", "expectedoutcome", "text"},
}

_HEADER_SCAN_ROWS = 15
_PARAGRAPH_ID = re.compile(r"^\s*\[?([A-Za-z][A-Za-z0-9]*(?:[-_.][A-Za-z0-9]+)*[-_. ]?\d+(?:\.\d+)*)\]?\s*[:\-–—)]\s*(.+)$")
_TABLE_SUFFIXES = {".xlsx", ".xlsm", ".csv", ".docx"}
_TEST_SUFFIXES = _TABLE_SUFFIXES | {".json", ".py"}


def load_requirements(path: str | Path) -> list[Requirement]:
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in _TABLE_SUFFIXES:
        raise ParseError(f"Unsupported requirements format '{suffix}'. Use .docx, .xlsx or .csv.")
    records = _table_records(path, REQUIREMENT_COLUMNS)
    if not records:
        raise ParseError(
            f"No requirements found in {path.name}. Expected a table with an ID column and a "
            "description column, or paragraphs such as 'REQ-001: The system shall ...'."
        )
    requirements = [Requirement(r["id"], r["text"] or r["title"], r["title"] if r["text"] else "", path.name)
                    for r in records]
    _check_unique((r.id for r in requirements), "requirement", path.name)
    return requirements


def load_tests(path: str | Path) -> list[TestCase]:
    path = Path(path)
    if path.is_dir():
        files = sorted(p for p in path.rglob("*") if p.is_file() and p.suffix.lower() in _TEST_SUFFIXES
                       and not p.name.startswith(("~$", ".")))
        tests = [t for f in files for t in _tests_from_file(f)]
        label = str(path)
    else:
        if path.suffix.lower() not in _TEST_SUFFIXES:
            raise ParseError(f"Unsupported test repository format '{path.suffix}'. "
                             "Use .xlsx, .csv, .docx, .json, .py or a directory of these.")
        tests = _tests_from_file(path)
        label = path.name
    if not tests:
        raise ParseError(f"No test cases found in {label}. Expected a table with an ID column and a "
                         "title, steps or description column.")
    _check_unique((t.id for t in tests), "test case", label)
    return tests


def _tests_from_file(path: Path) -> list[TestCase]:
    suffix = path.suffix.lower()
    if suffix == ".py":
        return _python_tests(path)
    if suffix == ".json":
        return _json_tests(path)
    return [TestCase(r["id"], r["text"], r["title"], path.name) for r in _table_records(path, TEST_COLUMNS)]


def _check_unique(ids: Iterable[str], kind: str, label: str) -> None:
    seen: set[str] = set()
    for item_id in ids:
        if item_id in seen:
            raise ParseError(f"Duplicate {kind} ID '{item_id}' in {label}.")
        seen.add(item_id)


# --- tabular sources ---------------------------------------------------------

def _table_records(path: Path, columns: dict[str, set[str]]) -> list[dict[str, str]]:
    suffix = path.suffix.lower()
    try:
        if suffix == ".csv":
            tables = [_csv_rows(path)]
        elif suffix == ".docx":
            tables, paragraphs = _docx_content(path)
        else:
            tables = _xlsx_tables(path)
    except ParseError:
        raise
    except Exception as exc:  # corrupt or mislabelled files surface as library-specific errors
        raise ParseError(f"Could not read {path.name}: {exc}") from exc

    records = [rec for rows in tables for rec in _records_from_rows(rows, columns)]
    if not records and suffix == ".docx":
        records = _records_from_paragraphs(paragraphs)
    return records


def _normalise(header: str) -> str:
    return re.sub(r"[^a-z0-9]", "", header.lower())


def _map_header(row: list[str], columns: dict[str, set[str]]) -> dict[str, list[int]] | None:
    mapping: dict[str, list[int]] = {"id": [], "title": [], "text": []}
    for i, cell in enumerate(row):
        key = _normalise(cell)
        for role, aliases in columns.items():
            if key in aliases:
                mapping[role].append(i)
                break
    if mapping["id"] and (mapping["text"] or mapping["title"]):
        return mapping
    return None


def _records_from_rows(rows: list[list[str]], columns: dict[str, set[str]]) -> list[dict[str, str]]:
    for header_at, row in enumerate(rows[:_HEADER_SCAN_ROWS]):
        mapping = _map_header(row, columns)
        if mapping:
            break
    else:
        return []

    def cell(row: list[str], i: int) -> str:
        return row[i].strip() if i < len(row) else ""

    records = []
    for row in rows[header_at + 1:]:
        item_id = cell(row, mapping["id"][0])
        title = " ".join(filter(None, (cell(row, i) for i in mapping["title"])))
        text = " ".join(filter(None, (cell(row, i) for i in mapping["text"])))
        if item_id and (title or text):
            records.append({"id": item_id, "title": title, "text": text})
    return records


def _records_from_paragraphs(paragraphs: list[str]) -> list[dict[str, str]]:
    records = []
    for paragraph in paragraphs:
        match = _PARAGRAPH_ID.match(paragraph)
        if match:
            records.append({"id": match.group(1).strip(), "title": "", "text": match.group(2).strip()})
    return records


def _cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _csv_rows(path: Path) -> list[list[str]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        return [[c.strip() for c in row] for row in csv.reader(handle)]


def _xlsx_tables(path: Path) -> list[list[list[str]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return [[[_cell_text(v) for v in row] for row in sheet.iter_rows(values_only=True)]
                for sheet in workbook.worksheets]
    finally:
        workbook.close()


def _docx_content(path: Path) -> tuple[list[list[list[str]]], list[str]]:
    from docx import Document
    from docx.oxml.ns import qn

    document = Document(str(path))
    # Walk the XML directly: python-docx's row.cells is quadratic on large tables.
    tables = [
        [["".join(t.text or "" for t in tc.iter(qn("w:t"))).strip() for tc in tr.iterchildren(qn("w:tc"))]
         for tr in table._tbl.iterchildren(qn("w:tr"))]
        for table in document.tables
    ]
    paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    return tables, paragraphs


# --- non-tabular test sources --------------------------------------------------

def _json_tests(path: Path) -> list[TestCase]:
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise ParseError(f"Could not read {path.name}: {exc}") from exc
    if isinstance(data, dict):
        data = next((v for v in data.values() if isinstance(v, list)), [])
    if not isinstance(data, list) or not all(isinstance(item, dict) for item in data):
        return []
    header = sorted({key for item in data for key in item})
    rows = [header] + [[_json_cell(item.get(key)) for key in header] for item in data]
    return [TestCase(r["id"], r["text"], r["title"], path.name) for r in _records_from_rows(rows, TEST_COLUMNS)]


def _json_cell(value: object) -> str:
    if isinstance(value, list):
        return " ".join(_cell_text(v) for v in value)
    return _cell_text(value)


def _python_tests(path: Path) -> list[TestCase]:
    try:
        tree = ast.parse(path.read_text(encoding="utf-8"))
    except (OSError, SyntaxError, UnicodeDecodeError) as exc:
        raise ParseError(f"Could not read {path.name}: {exc}") from exc

    tests: list[TestCase] = []

    def visit(nodes: list[ast.stmt], prefix: str) -> None:
        for node in nodes:
            if isinstance(node, ast.ClassDef) and node.name.startswith("Test"):
                visit(node.body, f"{prefix}{node.name}::")
            elif isinstance(node, (ast.FunctionDef, ast.AsyncFunctionDef)) and node.name.startswith("test"):
                title = node.name.removeprefix("test").replace("_", " ").strip()
                docstring = " ".join((ast.get_docstring(node) or "").split())
                tests.append(TestCase(f"{path.name}::{prefix}{node.name}", docstring, title, path.name))

    visit(tree.body, "")
    return tests
