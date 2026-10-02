"""Generate the bundled benchmark: 300 requirements, a test repository and labelled links.

The set is synthetic. Requirements and tests are worded differently on purpose: tests
describe behaviour instead of restating the requirement, and SYNONYM_SHARE of them also call
the business object by another name ("bill" for "invoice"), so matching needs semantics
rather than shared keywords. Labels are known by construction: 27 requirements have no test
and 14 tests belong to no requirement.

    python scripts/generate_sample_data.py [output_dir]
"""

from __future__ import annotations

import csv
import random
import sys
from pathlib import Path

from docx import Document
from openpyxl import Workbook

SEED = 7
UNCOVERED_REQUIREMENTS = 27
TWO_TEST_SHARE = 0.6
SYNONYM_SHARE = 0.25  # tests that use a different term from the requirement for the business object

# (requirement singular, requirement plural, test singular, test plural)
ENTITIES = [
    ("customer account", "customer accounts", "client profile", "client profiles"),
    ("invoice", "invoices", "bill", "bills"),
    ("purchase order", "purchase orders", "procurement request", "procurement requests"),
    ("support ticket", "support tickets", "helpdesk case", "helpdesk cases"),
    ("product listing", "product listings", "catalogue item", "catalogue items"),
    ("shipment", "shipments", "delivery consignment", "delivery consignments"),
    ("employee record", "employee records", "staff file", "staff files"),
    ("payment record", "payment records", "transaction entry", "transaction entries"),
    ("discount coupon", "discount coupons", "promo code", "promo codes"),
    ("supplier contract", "supplier contracts", "vendor agreement", "vendor agreements"),
    ("warehouse location", "warehouse locations", "storage site", "storage sites"),
    ("refund request", "refund requests", "money-back claim", "money-back claims"),
    ("appointment booking", "appointment bookings", "scheduled visit", "scheduled visits"),
    ("subscription plan", "subscription plans", "membership tier", "membership tiers"),
    ("expense claim", "expense claims", "reimbursement submission", "reimbursement submissions"),
]

# (requirement, [(test title, steps, expected result), (..)])
# {e}/{ep} = requirement wording, {s}/{sp} = test wording
OPERATIONS = [
    ("The system shall allow an authorised user to create a new {e} by entering all mandatory fields.",
     [("Add {s} with valid data",
       "Open the new {s} form, fill in every required field with valid values and save.",
       "The {s} is stored and appears in the list."),
      ("Add {s} with a required field missing",
       "Open the new {s} form, leave one required field blank and attempt to save.",
       "Saving is blocked and the blank field is highlighted.")]),
    ("The system shall allow an authorised user to modify the details of an existing {e}.",
     [("Update existing {s}",
       "Open a saved {s}, change several values and save.",
       "The changed values are shown when the {s} is reopened."),
      ("Discard update of {s}",
       "Open a saved {s}, change a value and then discard the changes.",
       "The original values of the {s} are unchanged.")]),
    ("The system shall require explicit confirmation before a {e} is deleted and shall remove it from "
     "active lists once confirmed.",
     [("Remove {s} after confirming",
       "Choose to remove a {s} and accept the prompt that follows.",
       "The {s} no longer appears in the active list."),
      ("Decline removal prompt for {s}",
       "Choose to remove a {s} and dismiss the prompt that follows.",
       "The {s} is still present.")]),
    ("The system shall provide a keyword search that returns every {e} whose name or reference "
     "contains the entered text.",
     [("Find {s} by partial name",
       "Type part of a known {s} name into the lookup box and submit.",
       "All {sp} containing the typed text are returned."),
      ("Find {s} with no hits",
       "Type a string that matches no {s} into the lookup box.",
       "An empty result message is displayed.")]),
    ("The system shall allow the list of {ep} to be filtered by status and by creation date range.",
     [("Narrow {s} list by state and dates",
       "Pick a state and a from/to date on the {s} list and apply.",
       "Only {sp} in that state created within the dates remain."),
      ("Reset narrowing on {s} list",
       "Apply a state and date selection to the {s} list, then reset it.",
       "The full list of {sp} is shown again.")]),
    ("The system shall export the currently displayed {ep} to a CSV file.",
     [("Download {s} list as comma-separated file",
       "Open the {s} list and choose the download option.",
       "A comma-separated file containing the listed {sp} is saved."),
      ("Download holds visible {sp} only",
       "Narrow the {s} list and then download it.",
       "The file holds only the {sp} that were visible on screen.")]),
    ("The system shall support bulk import of {ep} from an uploaded spreadsheet and shall reject rows "
     "that fail validation.",
     [("Load spreadsheet of {sp}",
       "Upload a well-formed spreadsheet containing several {sp}.",
       "Every row is loaded as a new {s}."),
      ("Load spreadsheet with bad rows of {sp}",
       "Upload a spreadsheet of {sp} in which some rows hold invalid values.",
       "The invalid rows are refused and listed in an error summary.")]),
    ("The system shall record an audit log entry containing the user, timestamp and changed values "
     "whenever a {e} is created, modified or deleted.",
     [("History entry written on {s} change",
       "Alter a {s} and open the activity history.",
       "An entry shows who made the change, when, and the old and new values."),
      ("History entry of {s} cannot be altered",
       "Attempt to edit or remove an activity history entry of a {s}.",
       "The entry is read-only.")]),
    ("The system shall restrict access to {ep} so that only users holding the appropriate role can "
     "view or change them.",
     [("User without permission blocked from {sp}",
       "Sign in as a user lacking the relevant permission and try to open a {s}.",
       "An access denied message is shown."),
      ("Permitted user can open {sp}",
       "Sign in with the relevant permission and open a {s}.",
       "The {s} is displayed and can be changed.")]),
    ("The system shall send an email notification to the owner whenever the status of a {e} changes.",
     [("Owner mailed when {s} state moves",
       "Move a {s} to a different state.",
       "The owner receives a message describing the new state."),
      ("No mail when {s} saved in same state",
       "Save a {s} without changing its state.",
       "No message is sent to the owner.")]),
    ("The system shall allow supporting documents of up to 10 MB each to be attached to a {e}.",
     [("Add file to {s}",
       "Add a 2 MB PDF to a {s}.",
       "The file is listed against the {s} and can be opened."),
      ("Oversized file refused on {s}",
       "Try to add a 15 MB file to a {s}.",
       "The upload is refused with a size limit message.")]),
    ("The system shall route each newly submitted {e} to a manager for approval before it becomes active.",
     [("Manager sign-off activates {s}",
       "Submit a {s} and have the manager sign it off.",
       "The {s} becomes active."),
      ("Manager turns down {s}",
       "Submit a {s} and have the manager turn it down with a reason.",
       "The {s} stays inactive and the submitter sees the reason.")]),
    ("The system shall detect a potential duplicate {e} at the time of entry and warn the user before saving.",
     [("Warning on entering identical {s}",
       "Enter a {s} whose key details equal those of one already stored.",
       "A warning that it may already exist appears before saving."),
      ("No warning for distinct {s}",
       "Enter a {s} whose details differ from all stored ones.",
       "It saves without any warning.")]),
    ("The system shall automatically archive any {e} that has been inactive for more than 24 months.",
     [("Old untouched {s} moved out of active list",
       "Run the nightly job with a {s} last touched over two years ago.",
       "The {s} is moved to long-term storage."),
      ("Recently used {s} kept",
       "Run the nightly job with a {s} touched last month.",
       "The {s} remains in the active list.")]),
    ("The system shall generate a printable PDF summary of a {e} on request.",
     [("Produce printable document for {s}",
       "Open a {s} and choose the print option.",
       "A document showing the {s} details is produced."),
      ("Printable document layout for {s}",
       "Produce the printable document for a {s} with long text values.",
       "All text is laid out within the page margins.")]),
    ("The system shall allow users to add time-stamped comments to a {e}.",
     [("Post note on {s}",
       "Open a {s} and post a note.",
       "The note is shown with the author name and the time it was posted."),
      ("Empty note refused on {s}",
       "Try to post a blank note on a {s}.",
       "Nothing is posted and a prompt asks for text.")]),
    ("The system shall retain previous versions of a {e} and allow an earlier version to be restored.",
     [("Roll {s} back to earlier revision",
       "Change a {s} twice, then pick the first revision and roll back to it.",
       "The {s} shows the values from the first revision."),
      ("Revision list for {s}",
       "Change a {s} several times and open its revision list.",
       "Each earlier revision is listed with its date.")]),
    ("The system shall display the list of {ep} within 2 seconds when up to 10,000 entries exist.",
     [("Load time of {s} list with large data",
       "Seed ten thousand {sp} and open the list while timing it.",
       "The list is rendered in under two seconds."),
      ("Paging through large {s} list",
       "With ten thousand {sp} seeded, move between pages of the list.",
       "Each page appears in under two seconds.")]),
    ("The system shall prevent two users from saving conflicting changes to the same {e} at the same time.",
     [("Simultaneous editing of one {s}",
       "Have two people open the same {s}; both change it and the second saves after the first.",
       "The second person is told the {s} was changed by someone else and nothing is overwritten."),
      ("Stale copy of {s} cannot be saved",
       "Keep a {s} open while a colleague saves it, then save the stale copy.",
       "The save is refused and a reload is offered.")]),
    ("The system shall validate date, numeric and email fields of a {e} and display a field-level "
     "error message for invalid input.",
     [("Malformed values rejected on {s} form",
       "Enter letters in a number box and a badly formed mail address on the {s} form.",
       "Each bad box shows its own error text."),
      ("Well-formed values accepted on {s} form",
       "Enter correct dates, numbers and mail addresses on the {s} form.",
       "No error text is shown.")]),
]

# Tests that trace to no requirement: half unrelated features, half familiar
# behaviours applied to business objects the requirements never mention.
ORPHAN_TESTS = [
    ("Dark theme toggle", "Switch the interface theme to dark in personal preferences.",
     "Every screen uses the dark colour scheme."),
    ("Legacy fax gateway sends document", "Queue an outbound fax through the legacy gateway.",
     "The transmission report shows success."),
    ("Password reset link expiry", "Request a password reset and open the link a day later.",
     "The link is reported as expired."),
    ("Interface language switch to French", "Choose French in the language menu.",
     "Menus and labels are shown in French."),
    ("Keyboard shortcut opens help panel", "Press F1 on the dashboard.",
     "The help panel slides in."),
    ("One-time code sent by text message at sign-in", "Sign in with two-step verification enabled.",
     "A six-digit code arrives by text message and is accepted."),
    ("Dashboard widget rearrangement", "Drag a chart widget to a new position on the home dashboard.",
     "The layout is remembered after signing out and back in."),
    ("Download loyalty points ledger as comma-separated file",
     "Open the loyalty points ledger and choose the download option.",
     "A comma-separated file containing the ledger lines is saved."),
    ("Add gift card with valid data",
     "Open the new gift card form, fill in every required field with valid values and save.",
     "The gift card is stored and appears in the list."),
    ("Owner mailed when warranty registration state moves",
     "Move a warranty registration to a different state.",
     "The owner receives a message describing the new state."),
    ("Add file to marketing campaign", "Add a 2 MB PDF to a marketing campaign.",
     "The file is listed against the marketing campaign and can be opened."),
    ("Find training course by partial name",
     "Type part of a known training course name into the lookup box and submit.",
     "All training courses containing the typed text are returned."),
    ("Post note on fleet vehicle", "Open a fleet vehicle and post a note.",
     "The note is shown with the author name and the time it was posted."),
    ("Manager sign-off activates timesheet", "Submit a timesheet and have the manager sign it off.",
     "The timesheet becomes active."),
]


def build() -> tuple[list[dict], list[dict], list[tuple[str, str]]]:
    rng = random.Random(SEED)

    requirements = []
    for e, ep, s, sp in ENTITIES:
        for template, test_templates in OPERATIONS:
            requirements.append({
                "id": f"REQ-{len(requirements) + 1:03d}",
                "module": f"{e.capitalize()} management",
                "text": template.format(e=e, ep=ep),
                "priority": rng.choice(["High", "Medium", "Low"]),
                "tests": test_templates,
                "terms": {True: (s, sp), False: (e, ep)},
            })

    uncovered = set(rng.sample(range(len(requirements)), UNCOVERED_REQUIREMENTS))
    tests = []
    for i, req in enumerate(requirements):
        if i in uncovered:
            continue
        count = 2 if rng.random() < TWO_TEST_SHARE else 1
        for template in req["tests"][:count]:
            s, sp = req["terms"][rng.random() < SYNONYM_SHARE]
            tests.append({"requirement_id": req["id"],
                          "fields": tuple(part.format(s=s, sp=sp) for part in template)})
    tests += [{"requirement_id": None, "fields": fields} for fields in ORPHAN_TESTS]

    rng.shuffle(tests)
    for i, test in enumerate(tests, start=1):
        test["id"] = f"TC-{i:04d}"
    truth = [(t["requirement_id"], t["id"]) for t in tests if t["requirement_id"]]
    return requirements, tests, sorted(truth)


def write_requirements_docx(requirements: list[dict], path: Path) -> None:
    document = Document()
    document.add_heading("Software Requirements Specification", level=0)
    document.add_paragraph("Order-to-cash back-office platform. Functional requirements are listed per module.")
    modules: dict[str, list[dict]] = {}
    for req in requirements:
        modules.setdefault(req["module"], []).append(req)
    for number, (module, rows) in enumerate(modules.items(), start=1):
        document.add_heading(f"3.{number} {module}", level=2)
        table = document.add_table(rows=len(rows) + 1, cols=3)
        table.style = "Table Grid"
        for cell, header in zip(table.rows[0].cells, ("Req ID", "Requirement Description", "Priority")):
            cell.text = header
        for row, req in zip(table.rows[1:], rows):
            for cell, value in zip(row.cells, (req["id"], req["text"], req["priority"])):
                cell.text = value
    document.save(str(path))


def write_requirements_xlsx(requirements: list[dict], path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Requirements"
    sheet.append(["Req ID", "Module", "Requirement Description", "Priority"])
    for req in requirements:
        sheet.append([req["id"], req["module"], req["text"], req["priority"]])
    workbook.save(path)


def write_tests_xlsx(tests: list[dict], path: Path) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Test Cases"
    sheet.append(["Test Case ID", "Title", "Steps", "Expected Result"])
    for test in tests:
        sheet.append([test["id"], *test["fields"]])
    workbook.save(path)


def write_ground_truth(truth: list[tuple[str, str]], path: Path) -> None:
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["requirement_id", "test_id"])
        writer.writerows(truth)


def main() -> None:
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parent.parent / "data" / "sample"
    out.mkdir(parents=True, exist_ok=True)
    requirements, tests, truth = build()
    write_requirements_docx(requirements, out / "requirements.docx")
    write_requirements_xlsx(requirements, out / "requirements.xlsx")
    write_tests_xlsx(tests, out / "test_cases.xlsx")
    write_ground_truth(truth, out / "ground_truth.csv")
    covered = len({r for r, _ in truth})
    print(f"{len(requirements)} requirements ({len(requirements) - covered} without tests), "
          f"{len(tests)} test cases ({len(tests) - len(truth)} orphans), {len(truth)} labelled links -> {out}")


if __name__ == "__main__":
    main()
