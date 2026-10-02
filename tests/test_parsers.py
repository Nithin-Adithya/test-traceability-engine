from __future__ import annotations

import json

import pytest
from docx import Document
from openpyxl import Workbook

from traceability.parsers import ParseError, load_requirements, load_tests


def _docx_table(path, header, rows):
    document = Document()
    document.add_heading("Spec", level=1)
    table = document.add_table(rows=1, cols=len(header))
    for cell, value in zip(table.rows[0].cells, header):
        cell.text = value
    for row in rows:
        for cell, value in zip(table.add_row().cells, row):
            cell.text = value
    document.save(str(path))


def _xlsx(path, rows):
    workbook = Workbook()
    for row in rows:
        workbook.active.append(row)
    workbook.save(path)


def test_requirements_from_docx_table(tmp_path):
    path = tmp_path / "reqs.docx"
    _docx_table(path, ["Req ID", "Requirement Description", "Priority"],
                [["REQ-1", "The system shall log in users.", "High"],
                 ["REQ-2", "The system shall log out users.", "Low"]])
    reqs = load_requirements(path)
    assert [(r.id, r.text) for r in reqs] == [("REQ-1", "The system shall log in users."),
                                              ("REQ-2", "The system shall log out users.")]


def test_requirements_from_docx_paragraphs(tmp_path):
    path = tmp_path / "reqs.docx"
    document = Document()
    document.add_paragraph("Introduction text that is not a requirement.")
    document.add_paragraph("REQ-001: The system shall encrypt data at rest.")
    document.add_paragraph("SRS_2.1 - The system shall back up nightly.")
    document.save(str(path))
    reqs = load_requirements(path)
    assert [r.id for r in reqs] == ["REQ-001", "SRS_2.1"]
    assert reqs[1].text == "The system shall back up nightly."


def test_requirements_from_xlsx_with_preamble_rows_and_title(tmp_path):
    path = tmp_path / "reqs.xlsx"
    _xlsx(path, [["Project X requirements"], [], ["ID", "Title", "Description"],
                 [1, "Login", "Users can sign in."], [None, None, None], [2, "Logout", "Users can sign out."]])
    reqs = load_requirements(path)
    assert [r.id for r in reqs] == ["1", "2"]
    assert reqs[0].embedding_text == "Login. Users can sign in."


def test_requirements_reject_duplicates_and_unknown_layouts(tmp_path):
    dup = tmp_path / "dup.xlsx"
    _xlsx(dup, [["ID", "Description"], ["R1", "a"], ["R1", "b"]])
    with pytest.raises(ParseError, match="Duplicate requirement ID 'R1'"):
        load_requirements(dup)

    no_header = tmp_path / "plain.xlsx"
    _xlsx(no_header, [["foo", "bar"], ["1", "2"]])
    with pytest.raises(ParseError, match="No requirements found"):
        load_requirements(no_header)

    with pytest.raises(ParseError, match="Unsupported"):
        load_requirements(tmp_path / "reqs.pdf")


def test_tests_from_xlsx_joins_steps_and_expected(tmp_path):
    path = tmp_path / "tests.xlsx"
    _xlsx(path, [["Test Case ID", "Title", "Steps", "Expected Result"],
                 ["TC-1", "Login works", "Enter credentials", "Dashboard shown"]])
    (test,) = load_tests(path)
    assert test.id == "TC-1"
    assert test.embedding_text == "Login works. Enter credentials Dashboard shown"


def test_tests_from_directory_mixes_formats(tmp_path):
    (tmp_path / "manual.csv").write_text("Test ID,Summary\nM-1,Check login\n", encoding="utf-8")
    (tmp_path / "api.json").write_text(json.dumps([{"id": "J-1", "name": "Check logout",
                                                    "steps": ["click", "confirm"]}]), encoding="utf-8")
    (tmp_path / "test_auth.py").write_text(
        "def helper():\n    pass\n\n"
        "def test_password_reset():\n    '''Reset link is mailed.'''\n\n"
        "class TestLock:\n    def test_account_locks(self):\n        pass\n",
        encoding="utf-8",
    )
    (tmp_path / "notes.txt").write_text("ignored", encoding="utf-8")
    tests = {t.id: t for t in load_tests(tmp_path)}
    assert set(tests) == {"M-1", "J-1", "test_auth.py::test_password_reset",
                          "test_auth.py::TestLock::test_account_locks"}
    assert tests["J-1"].embedding_text == "Check logout. click confirm"
    assert tests["test_auth.py::test_password_reset"].embedding_text == "password reset. Reset link is mailed."
    assert tests["test_auth.py::TestLock::test_account_locks"].embedding_text == "account locks"


def test_empty_test_repository_is_an_error(tmp_path):
    with pytest.raises(ParseError, match="No test cases found"):
        load_tests(tmp_path)
