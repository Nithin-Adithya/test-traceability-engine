from __future__ import annotations

import pytest
from openpyxl import load_workbook

from traceability.engine import TraceConfig, TraceabilityEngine, select_links
from traceability.evaluate import evaluate, load_ground_truth, sweep
from traceability.models import Requirement, TestCase
from traceability.parsers import ParseError
from traceability.report import write_report

REQUIREMENTS = [
    Requirement("R1", "user can log in with a password"),
    Requirement("R2", "invoice totals are exported to a csv file"),
    Requirement("R3", "an audit trail records every deletion"),
]
TESTS = [
    TestCase("T1", "log in with a valid password", "user log in"),
    TestCase("T2", "log in with a wrong password is refused"),
    TestCase("T3", "invoice totals exported to csv file match the screen"),
    TestCase("T4", "dark theme toggle switches colours"),
]
TRUTH = {("R1", "T1"), ("R1", "T2"), ("R2", "T3")}


@pytest.fixture
def result(embedder):
    return TraceabilityEngine(embedder, TraceConfig(threshold=0.3, margin=0.0)).trace(REQUIREMENTS, TESTS)


def test_links_gaps_and_orphans(result):
    assert {(l.requirement_id, l.test_id) for l in result.links} == TRUTH
    assert [row["requirement_id"] for row in result.uncovered()] == ["R3"]
    assert [row["test_id"] for row in result.orphans()] == ["T4"]

    summary = result.summary()
    assert (summary["covered_requirements"], summary["uncovered_requirements"], summary["orphan_tests"]) == (2, 1, 1)
    assert summary["coverage_pct"] == 66.7

    matrix = {row["requirement_id"]: row for row in result.matrix()}
    assert matrix["R1"]["test_count"] == 2
    assert matrix["R3"]["status"] == "Uncovered" and matrix["R3"]["best_score"] is None


def test_select_links_threshold_and_margin():
    candidates = {"T1": [("R1", 0.80), ("R2", 0.78), ("R3", 0.60)], "T2": [("R1", 0.40)]}
    assert {(l.requirement_id, l.test_id) for l in select_links(candidates, 0.5, 0.0)} == {("R1", "T1")}
    assert {(l.requirement_id, l.test_id) for l in select_links(candidates, 0.5, 0.05)} == {("R1", "T1"), ("R2", "T1")}
    # the margin never admits a candidate that is below the threshold
    assert {l.requirement_id for l in select_links(candidates, 0.79, 0.5)} == {"R1"}


def test_with_config_relinks_without_reembedding(result):
    strict = result.with_config(TraceConfig(threshold=0.99))
    assert strict.links == []
    assert len(strict.orphans()) == len(TESTS)
    assert len(result.links) == 3  # original untouched


def test_engine_rejects_empty_inputs_and_bad_config(embedder):
    with pytest.raises(ValueError):
        TraceabilityEngine(embedder).trace([], TESTS)
    with pytest.raises(ValueError):
        TraceabilityEngine(embedder).trace(REQUIREMENTS, [])
    with pytest.raises(ValueError):
        TraceConfig(threshold=1.5)


def test_evaluate_perfect_and_degraded(result):
    scores = evaluate(result, TRUTH)
    assert scores["match_accuracy"] == 1.0
    assert scores["links"]["f1"] == 1.0
    assert scores["uncovered_requirements"] == {"predicted": 1, "actual": 1, "correct": 1,
                                                "precision": 1.0, "recall": 1.0, "f1": 1.0}

    # Nothing linked: only the true orphan (T4) is judged correct.
    scores = evaluate(result.with_config(TraceConfig(threshold=0.99)), TRUTH)
    assert scores["match_accuracy"] == 0.25
    assert scores["links"]["recall"] == 0.0
    assert scores["orphan_tests"]["recall"] == 1.0 and scores["orphan_tests"]["precision"] == 0.25

    with pytest.raises(ParseError, match="missing from the inputs"):
        evaluate(result, {("R9", "T1")})


def test_sweep_and_ground_truth_loading(result, tmp_path):
    path = tmp_path / "truth.csv"
    path.write_text("requirement_id,test_id\nR1,T1\nR1,T2\nR2,T3\n", encoding="utf-8")
    truth = load_ground_truth(path)
    assert truth == TRUTH
    rows = sweep(result, truth, [0.3, 0.99], [0.0])
    assert [row["match_accuracy"] for row in rows] == [1.0, 0.25]


def test_report_workbook(result, tmp_path):
    path = write_report(result, tmp_path / "report.xlsx", evaluate(result, TRUTH))
    workbook = load_workbook(path)
    assert workbook.sheetnames == ["Summary", "Traceability Matrix", "Links", "Uncovered Requirements", "Orphan Tests"]
    matrix = list(workbook["Traceability Matrix"].iter_rows(values_only=True))
    assert matrix[0][:3] == ("Requirement ID", "Requirement", "Status")
    assert [row[2] for row in matrix[1:]] == ["Covered", "Covered", "Uncovered"]
    assert [row[0] for row in workbook["Orphan Tests"].iter_rows(min_row=2, values_only=True)] == ["T4"]
    assert workbook["Links"].max_row == 4
