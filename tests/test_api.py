from __future__ import annotations

from io import BytesIO

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from traceability.api import app, embedder_dependency

REQUIREMENTS_CSV = (
    "Req ID,Description\n"
    "R1,user can log in with a password\n"
    "R2,invoice totals are exported to a csv file\n"
    "R3,an audit trail records every deletion\n"
)
TESTS_CSV = (
    "Test ID,Title\n"
    "T1,user log in with a valid password\n"
    "T3,invoice totals exported to csv file match the screen\n"
)
MORE_TESTS_CSV = "Test ID,Title\nT4,dark theme toggle switches colours\n"
TRUTH_CSV = "requirement_id,test_id\nR1,T1\nR2,T3\n"


@pytest.fixture
def client(embedder):
    app.dependency_overrides[embedder_dependency] = lambda: embedder
    yield TestClient(app)
    app.dependency_overrides.clear()


def _files(truth: bool = False):
    files = [
        ("requirements", ("reqs.csv", REQUIREMENTS_CSV, "text/csv")),
        ("tests", ("manual.csv", TESTS_CSV, "text/csv")),
        ("tests", ("extra.csv", MORE_TESTS_CSV, "text/csv")),
    ]
    if truth:
        files.append(("ground_truth", ("truth.csv", TRUTH_CSV, "text/csv")))
    return files


def test_health(client):
    assert client.get("/health").json() == {"status": "ok"}


def test_trace_returns_matrix_gaps_and_evaluation(client):
    response = client.post("/trace", files=_files(truth=True), data={"threshold": "0.3", "margin": "0"})
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["summary"]["requirements"] == 3 and body["summary"]["tests"] == 3
    assert [row["requirement_id"] for row in body["uncovered"]] == ["R3"]
    assert [row["test_id"] for row in body["orphans"]] == ["T4"]
    assert body["evaluation"]["match_accuracy"] == 1.0

    stored = client.get(f"/runs/{body['run_id']}").json()
    assert stored["links"] == body["links"]

    report = client.get(f"/runs/{body['run_id']}/report.xlsx")
    assert report.status_code == 200
    assert "Traceability Matrix" in load_workbook(BytesIO(report.content)).sheetnames


def test_trace_without_ground_truth(client):
    body = client.post("/trace", files=_files(), data={"threshold": "0.3"}).json()
    assert body["evaluation"] is None


def test_bad_inputs_are_422(client):
    files = [("requirements", ("reqs.txt", "nothing", "text/plain")),
             ("tests", ("manual.csv", TESTS_CSV, "text/csv"))]
    response = client.post("/trace", files=files)
    assert response.status_code == 422
    assert "Unsupported requirements format" in response.json()["detail"]

    response = client.post("/trace", files=_files(), data={"threshold": "7"})
    assert response.status_code == 422


def test_unknown_run_is_404(client):
    assert client.get("/runs/nope").status_code == 404
    assert client.get("/runs/nope/report.xlsx").status_code == 404
